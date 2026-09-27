#!/usr/bin/env python3
"""İndirilen resmi ek dosyalarından metin çıkar (PDF/Word/Excel/görsel OCR).

Her dosya için yanına <dosya>.txt yazar. OCR satırları konumlarına göre (y, x) sıralanır;
aynı satırdaki hücreler " | " ile ayrılır.
"""

from __future__ import annotations

import io
import sys
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FILES = ROOT / "data_out" / "fiyat" / "files"

_ocr = None


def ocr_image_bytes(data: bytes) -> str:
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    res, _ = _ocr(data)
    if not res:
        return ""
    items = []
    for box, text, score in res:
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        items.append(((min(ys) + max(ys)) / 2, min(xs), max(ys) - min(ys), text))
    items.sort()
    lines, cur, cur_y, cur_h = [], [], None, 0
    for y, x, h, t in items:
        if cur_y is None or abs(y - cur_y) <= max(8, 0.5 * max(h, cur_h)):
            cur.append((x, t))
            cur_y = y if cur_y is None else (cur_y + y) / 2
            cur_h = max(cur_h, h)
        else:
            lines.append(" | ".join(t for _, t in sorted(cur)))
            cur, cur_y, cur_h = [(x, t)], y, h
    if cur:
        lines.append(" | ".join(t for _, t in sorted(cur)))
    return "\n".join(lines)


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    data = p.read_bytes()
    if ext == ".pdf":
        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        out = []
        for i, page in enumerate(doc):
            txt = page.get_text("text")
            if len(txt.strip()) < 40:
                pix = page.get_pixmap(dpi=200)
                txt = "[OCR]\n" + ocr_image_bytes(pix.tobytes("png"))
            out.append(f"=== SAYFA {i + 1}\n{txt}")
            for t in page.find_tables().tables:
                out.append("--- TABLO")
                for row in t.extract():
                    out.append(" | ".join((c or "").replace("\n", " ") for c in row))
        return "\n".join(out)
    if ext in (".jpg", ".jpeg", ".png", ".webp"):
        return "[OCR]\n" + ocr_image_bytes(data)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        out = [para.text for para in d.paragraphs if para.text.strip()]
        for t in d.tables:
            out.append("--- TABLO")
            for row in t.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(out)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"=== SHEET {ws.title} (gizli={ws.sheet_state != 'visible'})")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        out = []
        for sh in wb.sheets():
            out.append(f"=== SHEET {sh.name}")
            for r in range(sh.nrows):
                out.append(" | ".join(str(v) for v in sh.row_values(r)))
        return "\n".join(out)
    if ext == ".doc":
        return "[DOC eski biçim — metin çıkarılamadı]"
    return ""


def work(p: Path) -> tuple[str, int]:
    out = p.with_suffix(p.suffix + ".txt")
    if out.is_file():
        return p.name, -1
    try:
        txt = extract(p)
    except Exception as e:  # noqa: BLE001
        txt = f"[HATA] {type(e).__name__}: {e}"
    out.write_text(txt, encoding="utf-8")
    return p.name, len(txt)


def main() -> None:
    files = [p for p in FILES.iterdir() if p.suffix.lower() in (".pdf", ".jpg", ".jpeg", ".png", ".webp", ".docx", ".xlsx", ".xls", ".doc")]
    print("dosya", len(files), flush=True)
    with ProcessPoolExecutor(max_workers=4) as ex:
        for i, (name, n) in enumerate(ex.map(work, files), 1):
            if i % 50 == 0:
                print(f"  {i}/{len(files)}", flush=True)
    print("bitti")


if __name__ == "__main__":
    sys.exit(main())
